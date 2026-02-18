"""
Tests for user document upload processor.
Covers PDF/CSV/DOCX/TXT routing, DOCX extraction, TXT chunking, and missing FIPS handling.
"""

import io
from unittest.mock import MagicMock, patch

import pytest

from data.schemas.document import ContentType, CountyDocument, DocumentCategory
from pipelines.ingest.upload_processor import (
    CHUNK_SIZE,
    OVERLAP,
    USER_UPLOAD_COUNTY,
    USER_UPLOAD_FIPS,
    USER_UPLOAD_STATE,
    UploadProcessor,
    _chunk_text,
    _detect_file_type,
)


# ---- Helpers ----


def _make_pdf_bytes(text: str = "Sample PDF content for county records.") -> bytes:
    """Minimal PDF bytes (magic + minimal body)."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text[:5000])
    buf = io.BytesIO()
    doc.save(buf, deflate=True)
    doc.close()
    buf.seek(0)
    return buf.read()


def _make_csv_bytes(content: str = "col1,col2,col3\na,b,c\n1,2,3") -> bytes:
    """CSV file bytes."""
    return content.encode("utf-8")


def _make_docx_bytes_with_text(text: str) -> bytes:
    """Create minimal DOCX with given text using python-docx."""
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ---- _detect_file_type ----


class TestDetectFileType:
    """Test file type detection by magic bytes and extension."""

    def test_pdf_by_magic(self) -> None:
        assert _detect_file_type(b"%PDF-1.4\n", "x.bin") == "pdf"

    def test_pdf_by_extension(self) -> None:
        assert _detect_file_type(b"not really pdf", "file.pdf") == "pdf"

    def test_docx_by_magic_and_extension(self) -> None:
        # DOCX starts with PK (zip)
        assert _detect_file_type(b"PK\x03\x04", "doc.docx") == "docx"

    def test_csv_by_extension(self) -> None:
        assert _detect_file_type(b"a,b,c\n1,2,3", "data.csv") == "csv"

    def test_txt_by_extension(self) -> None:
        assert _detect_file_type(b"hello world", "readme.txt") == "txt"

    def test_md_by_extension(self) -> None:
        assert _detect_file_type(b"# Title", "notes.md") == "md"

    def test_unknown_extension(self) -> None:
        assert _detect_file_type(b"binary", "file.xyz") == "unknown"


# ---- _chunk_text ----


class TestChunkText:
    """Test text chunking (same as PDFProcessor)."""

    def test_chunk_text_empty_returns_empty_list(self) -> None:
        assert _chunk_text("") == []
        assert _chunk_text("   \n  ") == []

    def test_chunk_text_short_returns_one_chunk(self) -> None:
        text = "Short piece of text."
        assert len(_chunk_text(text)) == 1
        assert _chunk_text(text)[0] == text

    def test_chunk_text_long_produces_multiple_chunks(self) -> None:
        text = "x" * (CHUNK_SIZE * 3)
        chunks = _chunk_text(text)
        assert len(chunks) >= 2
        assert sum(len(c) for c in chunks) >= len(text) - OVERLAP * (len(chunks) - 1)


# ---- UploadProcessor - init ----


class TestUploadProcessorInit:
    """Test UploadProcessor initialization."""

    def test_default_chunk_and_overlap(self) -> None:
        p = UploadProcessor()
        assert p.chunk_size == CHUNK_SIZE
        assert p.overlap == OVERLAP

    def test_custom_chunk_overlap(self) -> None:
        p = UploadProcessor(chunk_size=500, overlap=50)
        assert p.chunk_size == 500
        assert p.overlap == 50


# ---- PDF routing ----


class TestUploadProcessorPDF:
    """Test that PDFs are routed to PDFProcessor."""

    def test_pdf_returns_county_documents(self) -> None:
        processor = UploadProcessor()
        pdf_bytes = _make_pdf_bytes("County deed 2024. Parcel 12345.")
        docs = processor.process(pdf_bytes, "deed.pdf", fips="01001", category=DocumentCategory.legal)
        assert isinstance(docs, list)
        assert len(docs) >= 1
        for d in docs:
            assert isinstance(d, CountyDocument)
            assert d.content_type == ContentType.pdf or d.content_type == ContentType.text
            assert d.metadata.get("file_type") == "pdf"
            assert d.metadata.get("upload_source") == "user"
            assert d.metadata.get("original_filename") == "deed.pdf"

    def test_pdf_without_fips_uses_placeholder_fips(self) -> None:
        processor = UploadProcessor()
        pdf_bytes = _make_pdf_bytes("User upload content.")
        docs = processor.process(pdf_bytes, "up.pdf", fips=None, category=DocumentCategory.user_upload)
        assert len(docs) >= 1
        assert docs[0].fips == USER_UPLOAD_FIPS
        assert docs[0].county_name == USER_UPLOAD_COUNTY
        assert docs[0].state_abbr == USER_UPLOAD_STATE


# ---- CSV routing ----


class TestUploadProcessorCSV:
    """Test that CSV is routed to StructuredProcessor."""

    def test_csv_returns_structured_documents(self) -> None:
        processor = UploadProcessor()
        csv_bytes = _make_csv_bytes("name,value\nAlice,100\nBob,200")
        docs = processor.process(csv_bytes, "data.csv", fips="01001", category=DocumentCategory.property)
        assert isinstance(docs, list)
        # One doc per row (excluding header)
        assert len(docs) >= 1
        for d in docs:
            assert isinstance(d, CountyDocument)
            assert d.content_type == ContentType.structured
            assert d.metadata.get("file_type") == "csv"
            assert d.metadata.get("upload_source") == "user"

    def test_csv_without_fips_uses_placeholder(self) -> None:
        processor = UploadProcessor()
        csv_bytes = _make_csv_bytes("a,b\n1,2")
        docs = processor.process(csv_bytes, "data.csv", fips=None)
        assert len(docs) >= 1
        assert docs[0].fips == USER_UPLOAD_FIPS
        assert docs[0].county_name == USER_UPLOAD_COUNTY
        assert docs[0].state_abbr == USER_UPLOAD_STATE


# ---- DOCX ----


class TestUploadProcessorDOCX:
    """Test DOCX text extraction and chunking."""

    def test_extract_docx_returns_text(self) -> None:
        processor = UploadProcessor()
        text = "This is a Word document with county information."
        docx_bytes = _make_docx_bytes_with_text(text)
        out = processor.extract_docx(docx_bytes)
        assert "county" in out.lower() or "Word" in out or "document" in out

    def test_process_docx_produces_documents(self) -> None:
        processor = UploadProcessor()
        long_text = "County report. " + ("Detail line. " * 80)
        docx_bytes = _make_docx_bytes_with_text(long_text)
        docs = processor.process(docx_bytes, "report.docx", fips=None, category=DocumentCategory.user_upload)
        assert len(docs) >= 1
        for d in docs:
            assert d.metadata.get("file_type") == "docx"
            assert d.metadata.get("upload_source") == "user"
            assert "chunk_index" in d.metadata
            assert "total_chunks" in d.metadata


# ---- TXT / MD ----


class TestUploadProcessorTXT:
    """Test TXT/MD handling and chunking."""

    def test_extract_text_utf8(self) -> None:
        processor = UploadProcessor()
        raw = "Hello world. UTF-8 content."
        out = processor.extract_text(raw.encode("utf-8"), "f.txt")
        assert out == raw

    def test_process_txt_chunking(self) -> None:
        processor = UploadProcessor(chunk_size=100, overlap=20)
        text = "A" * 250
        txt_bytes = text.encode("utf-8")
        docs = processor.process(txt_bytes, "notes.txt", fips="01001", category=DocumentCategory.user_upload)
        assert len(docs) >= 2
        for d in docs:
            assert d.metadata.get("file_type") == "txt"
            assert d.metadata.get("upload_source") == "user"
            assert d.metadata.get("chunk_index") is not None
            assert d.metadata.get("total_chunks") == len(docs)

    def test_process_md_single_chunk(self) -> None:
        processor = UploadProcessor()
        content = "# Title\n\nShort markdown body."
        docs = processor.process(content.encode("utf-8"), "readme.md", fips=None)
        assert len(docs) == 1
        assert "Title" in docs[0].raw_content or "markdown" in docs[0].raw_content


# ---- Missing FIPS ----


class TestUploadProcessorMissingFips:
    """Test that missing FIPS is handled gracefully."""

    def test_none_fips_uses_00000_and_placeholders(self) -> None:
        processor = UploadProcessor()
        txt = "User upload with no county." + " x" * 30
        docs = processor.process(txt.encode("utf-8"), "u.txt", fips=None)
        assert len(docs) >= 1
        assert docs[0].fips == USER_UPLOAD_FIPS
        assert docs[0].county_name == USER_UPLOAD_COUNTY
        assert docs[0].state_abbr == USER_UPLOAD_STATE

    def test_empty_string_fips_treated_as_none(self) -> None:
        processor = UploadProcessor()
        txt = "Content. " * 20
        docs = processor.process(txt.encode("utf-8"), "u.txt", fips="", category=DocumentCategory.user_upload)
        assert len(docs) >= 1
        assert docs[0].fips == USER_UPLOAD_FIPS


# ---- Unsupported type ----


class TestUploadProcessorUnsupported:
    """Test unsupported file type returns empty list."""

    def test_unknown_type_returns_empty_list(self) -> None:
        processor = UploadProcessor()
        docs = processor.process(b"binary content", "file.xyz", fips="01001")
        assert docs == []
